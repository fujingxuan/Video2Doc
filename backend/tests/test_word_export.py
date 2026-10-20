from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument

from app.core.errors import PipelineError
from app.documents.exporter import ExportConfig, export_document_bundle
from app.documents.word import render_word_document
from tests.test_documents import sample_document


def test_word_retains_final_text_formatting_images_and_source(tmp_path):
    document = sample_document(tmp_path)
    document.processing.update(article_mode="full_article", keep_timestamps=False)
    block = document.all_blocks()[0]
    block.edited_text = (
        "### 校对后的正文\n\n保留数字 **42** 和 *APIName*。\n\n"
        "1. 第一项\n2. 第二项\n\n"
        "| 内容 | 数值 |\n| --- | --- |\n| 质量 | 42 |\n\n"
        "```python\nprint(42)\n```\n\n[参考来源](https://example.com)"
    )
    result = export_document_bundle(document, ExportConfig(vault_path=str(tmp_path)))
    word = WordDocument(result.word_path)
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "校对后的正文" in text and "保留数字 42 和 APIName。" in text
    assert block.original_text not in text
    assert "第二段继续原始论证，不生成总结。" in text
    assert "1. 第一项" in text and "2. 第二项" in text and "print(42)" in text
    assert document.source.source in text and "00:00:10" not in text
    assert word.paragraphs[0].style.name == "Title"
    assert word.tables[0].cell(1, 1).text == "42"
    with ZipFile(result.word_path) as archive:
        assert archive.testzip() is None
        assert any(name.startswith("word/media/") for name in archive.namelist())
        assert b"https://example.com" in archive.read("word/_rels/document.xml.rels")


@pytest.mark.parametrize("policy", ["rename", "fail"])
def test_word_alone_causes_bundle_filename_conflict(tmp_path, policy):
    document = sample_document(tmp_path)
    config = ExportConfig(vault_path=str(tmp_path), filename_template="note",
                          inbox_relative_path=".", conflict_policy=policy)
    original = tmp_path / "note.docx"
    original.write_bytes(b"keep user file")
    if policy == "fail":
        with pytest.raises(PipelineError, match="已经存在"):
            export_document_bundle(document, config)
    else:
        result = export_document_bundle(document, config)
        assert Path(result.word_path).name == "note (2).docx"
        assert Path(result.markdown_path).name == "note (2).md"
    assert original.read_bytes() == b"keep user file"


def test_partial_commit_restores_both_formats_on_overwrite_failure(tmp_path, monkeypatch):
    document = sample_document(tmp_path)
    config = ExportConfig(vault_path=str(tmp_path), inbox_relative_path=".",
                          filename_template="note", conflict_policy="overwrite")
    initial = export_document_bundle(document, config)
    contents = {Path(path): Path(path).read_bytes()
                for path in (initial.markdown_path, initial.word_path)}
    original_replace = Path.replace

    def fail_markdown_commit(path, target):
        if path.name == "note.md" and path.parent.name.startswith(".video2doc-export-"):
            assert Path(initial.word_path).is_file()
            raise OSError("controlled commit failure")
        return original_replace(path, target)

    monkeypatch.setattr(Path, "replace", fail_markdown_commit)
    with pytest.raises(OSError, match="controlled commit failure"):
        export_document_bundle(document, config)
    assert all(path.read_bytes() == content for path, content in contents.items())
    assert not list(tmp_path.glob(".video2doc-*.tmp"))


def test_word_does_not_read_images_outside_export_directory(tmp_path):
    from PIL import Image
    outside = tmp_path / "private.png"
    Image.new("RGB", (20, 20), "red").save(outside)
    root = tmp_path / "export"
    root.mkdir()
    path = root / "note.docx"
    render_word_document("# 标题\n\n![外部图片](../private.png)", path, asset_root=root)
    with ZipFile(path) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())
    assert "../private.png" in "\n".join(p.text for p in WordDocument(path).paragraphs)

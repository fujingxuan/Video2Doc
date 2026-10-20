"""Render the exported Markdown as an editable Word document with embedded images."""
from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt

from app.core.errors import PipelineError


def render_word_document(markdown: str, destination: Path, *, asset_root: Path) -> None:
    """Use the same visible content as Markdown, excluding YAML and wrapper markup."""
    body = re.sub(r"\A---\r?\n.*?\r?\n---\r?\n", "", markdown, count=1, flags=re.S)
    body = re.sub(r"<!-- video2doc-sources:.*?-->", "", body)
    body = body.replace("<details>", "").replace("</details>", "")
    body = body.replace("<summary>处理与追溯信息</summary>", "## 处理与追溯信息")
    # Existing Obsidian image paths can contain spaces and parentheses.
    body = re.sub(r"^(!\[[^\n]*\])\((.+)\)$",
                  lambda match: f"{match[1]}({quote(match[2], safe='/')})",
                  body, flags=re.M)
    tokens = MarkdownIt("commonmark").enable("table").parse(body)
    word = WordDocument()
    section = word.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = word.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
    word.styles["Normal"].font.size = Pt(11)
    word.styles["Normal"].paragraph_format.space_after = Pt(6)
    word.styles["Normal"].paragraph_format.line_spacing = 1.25
    word.styles["Title"].font.size = Pt(22)
    word.styles["Title"].font.bold = True
    # Built-in Title styles can inherit an ornamental paragraph border.
    for name in ("Normal", "Title"):
        for border in word.styles[name].element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    word.core_properties.title = next(
        (token.content for token in tokens if token.type == "inline"), "Video2Doc"
    )
    word.core_properties.author = "Video2Doc"

    paragraph = None
    lists: list[dict] = []
    quote_depth = 0
    table = None
    row = None
    cell_index = 0
    first_title = True
    for index, token in enumerate(tokens):
        kind = token.type
        if kind == "heading_open":
            level = int(token.tag[1:])
            style = "Title" if first_title and level == 1 else f"Heading {min(level, 4)}"
            paragraph = word.add_paragraph(style=style)
            first_title = False
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            lists.append({"ordered": kind == "ordered_list_open",
                          "number": int(token.attrGet("start") or 1), "first": False})
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "list_item_open":
            lists[-1]["first"] = True
        elif kind == "blockquote_open":
            quote_depth += 1
        elif kind == "blockquote_close":
            quote_depth -= 1
        elif kind == "paragraph_open":
            paragraph = word.add_paragraph()
            if lists:
                paragraph.paragraph_format.left_indent = Cm(0.6 * len(lists))
                if lists[-1]["first"]:
                    prefix = f"{lists[-1]['number']}. " if lists[-1]["ordered"] else "• "
                    paragraph.add_run(prefix)
                    lists[-1]["number"] += 1
                    lists[-1]["first"] = False
            elif quote_depth:
                paragraph.paragraph_format.left_indent = Cm(0.6 * quote_depth)
        elif kind == "table_open":
            # Count the header cells before the first row closes.
            cells = 0
            for following in tokens[index + 1:]:
                if following.type == "tr_close":
                    break
                cells += following.type == "th_open"
            table = word.add_table(rows=0, cols=max(1, cells))
            table.style = "Table Grid"
            table.autofit = False
        elif kind == "tr_open":
            assert table is not None
            row = table.add_row()
            cell_index = 0
        elif kind in {"th_open", "td_open"}:
            assert row is not None
            paragraph = row.cells[cell_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            if kind == "th_open":
                paragraph.style = word.styles["Normal"]
            cell_index += 1
        elif kind == "thead_close" and row is not None:
            for cell in row.cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        elif kind == "table_close":
            table = row = None
            word.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "inline":
            if paragraph is None:
                paragraph = word.add_paragraph()
            _write_inline(paragraph, token.children or [], asset_root)
        elif kind in {"fence", "code_block"}:
            paragraph = word.add_paragraph()
            run = paragraph.add_run(token.content.rstrip("\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif kind in {"paragraph_close", "heading_close", "th_close", "td_close"}:
            paragraph = None
        # Horizontal rules and generated HTML wrappers carry no source content.
    destination.parent.mkdir(parents=True, exist_ok=True)
    word.save(destination)


def _write_inline(paragraph, tokens, asset_root: Path) -> None:
    bold = italic = 0
    hyperlink = None
    for token in tokens:
        kind = token.type
        if kind == "strong_open":
            bold += 1
        elif kind == "strong_close":
            bold -= 1
        elif kind == "em_open":
            italic += 1
        elif kind == "em_close":
            italic -= 1
        elif kind == "link_open":
            href = token.attrGet("href") or ""
            if urlparse(href).scheme in {"http", "https", "mailto"}:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, True))
                paragraph._p.append(hyperlink)
        elif kind == "link_close":
            hyperlink = None
        elif kind == "image":
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
            source = unquote(token.attrGet("src") or "")
            root = asset_root.resolve()
            image = (root / source).resolve()
            if not urlparse(source).scheme and image.is_relative_to(root) and image.is_file():
                run = paragraph.add_run()
                try:
                    picture = run.add_picture(str(image), width=Cm(16))
                except Exception as exc:
                    raise PipelineError("WORD_IMAGE_INVALID", "Word 无法嵌入相关画面",
                                        technical_detail=str(image)) from exc
                picture._inline.docPr.set("descr", token.content)
                if picture.height > Cm(21):
                    picture.width = int(picture.width * Cm(21) / picture.height)
                    picture.height = Cm(21)
            else:
                # Do not fetch external images or read outside the export staging area.
                paragraph.add_run(f"[图片：{token.content}]({source})")
        elif kind in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif kind in {"text", "code_inline", "html_inline"}:
            text = token.content
            if kind == "html_inline" and text.lower() in {"<br>", "<br/>", "<br />"}:
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(text)
            run.bold, run.italic = bool(bold), bool(italic)
            if kind == "code_inline":
                run.font.name = "Consolas"
            if hyperlink is not None:
                hyperlink.append(run._r)
